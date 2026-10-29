import json, shutil, sys
from datetime import datetime
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
MID = '0013nA7W2nTBBL'
LINK = (f'https://y.qq.com/m/client/music_index/index.html?ADTAG=cbshare&channelId=10036163'
        f'&mid={MID}&openinqqmusic=1&type={MID}')

# 1) albums.json 补 mid
pa = Path(r'D:\wx409.github.io\data\albums.json')
a = json.loads(pa.read_text(encoding='utf-8'))
for al in a['albums']:
    if al.get('name') == '火光':
        for s in al['songs']:
            s['mid'] = MID
pa.write_text(json.dumps(a, ensure_ascii=False, indent=1), encoding='utf-8')
print('albums.json 火光已补 mid')

# 2) 收听人数 xlsx 追加一行
px = Path(r'E:\wx\收听人数（2026.7.24）.xlsx')
bak = px.with_name(px.stem + f'.bak_{datetime.now():%Y%m%d_%H%M}' + px.suffix)
shutil.copy2(px, bak)
wb = openpyxl.load_workbook(px)
ws = wb.active
names = [str(ws.cell(row=r, column=2).value or '') for r in range(2, ws.max_row + 1)]
if any('火光' in n for n in names):
    print('已在清单中')
else:
    seq = ws.max_row  # 现有行数=序号
    r = ws.max_row + 1
    ws.cell(row=r, column=1).value = seq
    ws.cell(row=r, column=2).value = '火光'
    ws.cell(row=r, column=3).value = '王晰'
    ws.cell(row=r, column=4).value = LINK
    ws.cell(row=r, column=5).value = LINK
    ws.cell(row=r, column=7).value = '火光'
    wb.save(px)
    print('xlsx 已追加 火光（行', r, '，备份', bak.name, '）')
print('清单总行', ws.max_row)
