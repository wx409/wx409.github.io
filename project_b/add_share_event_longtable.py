import shutil, sys
from datetime import datetime
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
P = Path(r'E:\wx\index_records\历次巡演歌单\王晰巡演歌单长表_单一事实源.xlsx')
bak = P.with_name(P.stem + f'.bak_{datetime.now():%Y%m%d_%H%M}' + P.suffix)
shutil.copy2(P, bak)
print('备份:', bak.name)

wb = openpyxl.load_workbook(P)
ws = wb['合并长表']
rows = list(ws.iter_rows(min_row=2, values_only=True))
exists = any(str(r[1])[:10] == '2021-12-18' and r[2] and '重庆' in str(r[2]) for r in rows)
if exists:
    print('长表已有该场次')
else:
    ws.append([
        '新歌分享会', '2021-12-18', '重庆', None, None,
        '「万物晰有」新歌分享会，19:30；现场发布专辑《B面图景》；无公开歌单，曲目待补',
        None, '重庆发布 2021-12-16 / 小红书物料',
    ])
    wb.save(P)
    print('长表已追加：2021-12-18 重庆 新歌分享会')
print('长表现有行数', ws.max_row)
