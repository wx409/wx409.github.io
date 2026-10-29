import shutil, sys
from datetime import datetime
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
PX = Path(r'E:\wx\收听人数（2026.7.24）.xlsx')
MID = '0013nA7W2nTBBL'
LINK = (f'https://y.qq.com/m/client/music_index/index.html?ADTAG=cbshare&channelId=10036163'
        f'&mid={MID}&openinqqmusic=1&type={MID}')

bak = PX.with_name(PX.stem + f'.bak_{datetime.now():%Y%m%d_%H%M}' + PX.suffix)
shutil.copy2(PX, bak)
wb = openpyxl.load_workbook(PX)
ws = wb.active
names = [str(ws.cell(row=r, column=2).value or '') for r in range(2, ws.max_row + 1)]
if any('火光' in n for n in names):
    print('已存在火光')
else:
    r = ws.max_row + 1
    ws.cell(row=r, column=1).value = ws.max_row - 1
    ws.cell(row=r, column=2).value = '火光'
    ws.cell(row=r, column=3).value = '王晰'
    ws.cell(row=r, column=4).value = LINK
    ws.cell(row=r, column=5).value = LINK
    ws.cell(row=r, column=7).value = '火光'
    wb.save(PX)
    print('已写入主文件，行号', r, '｜备份', bak.name)
print('总行数', ws.max_row)
