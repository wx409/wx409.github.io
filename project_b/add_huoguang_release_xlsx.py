import shutil, sys
from datetime import datetime
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
P = Path(r'E:\wx\index_records\王晰歌曲信息汇总.xlsx')
wb = openpyxl.load_workbook(P)
ws = wb['OST&单曲']
names = [str(ws.cell(row=r, column=3).value or '').strip() for r in range(2, ws.max_row + 1)]
if '火光' in names:
    print('已存在火光')
else:
    bak = P.with_name(P.stem + f'.bak_{datetime.now():%Y%m%d_%H%M}' + P.suffix)
    shutil.copy2(P, bak)
    r = ws.max_row + 1
    ws.cell(row=r, column=1).value = '=ROW()-1'
    ws.cell(row=r, column=3).value = '火光'
    ws.cell(row=r, column=4).value = '2026-10-27'
    ws.cell(row=r, column=5).value = '单曲｜Columbia Records China X 韶愔音乐娱乐'
    wb.save(P)
    print('已写入 OST&单曲：火光 2026-10-27（备份', bak.name, '）')
print('行数', ws.max_row)
