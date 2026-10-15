# -*- coding: utf-8 -*-
"""长表行 204「演唱曲目」按**用户提供权威歌单图**补齐与改名（不删旧值：备注留痕）。

补：天上的风（傲日其愣独唱）、生日歌（加唱·双人）
改名：唱起草原歌 → 唱起草原的歌；晚安我心爱的姑娘 → 心爱的姑娘；岩缝中开出的花 → 岩峰里盛开的花
用法：python -X utf8 project_b/fix_row204_by_image.py [--apply]
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

XLSX = Path(r"E:\wx\index_records\王晰演出活动.xlsx")
BAK = XLSX.with_name("王晰演出活动.备份20261014_图源歌单.xlsx")
ROW = 204
RENAMES = {"唱起草原歌": "唱起草原的歌", "晚安我心爱的姑娘": "心爱的姑娘",
           "岩缝中开出的花": "岩峰里盛开的花"}
ADD = ["天上的风", "生日歌"]
MARK = "【图源补录 20261014】"


def main() -> int:
    apply = "--apply" in sys.argv
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    seq, note_c, song_c = hdr.index("序号") + 1, hdr.index("备注") + 1, hdr.index("演唱曲目") + 1
    r = next((i for i in range(2, ws.max_row + 1) if ws.cell(i, seq).value == ROW), None)
    if r is None:
        print("[X] 找不到行", ROW)
        return 1
    cur = [x.strip() for x in str(ws.cell(r, song_c).value or "").split("、") if x.strip()]
    n0 = len(cur)
    out, renamed = [], []
    for s in cur:
        if s in RENAMES:
            out.append(RENAMES[s])
            renamed.append(f"{s}→{RENAMES[s]}")
        else:
            out.append(s)
    added = []
    for s in ADD:
        if s not in out:
            out.append(s)
            added.append(s)
    note_add = (MARK + f"按用户提供权威歌单图补录：新增 {'、'.join(ADD) if ADD else '—'}（天上的风＝傲日其愣独唱；"
                f"生日歌＝加唱·双人）；曲名归一：{'；'.join(renamed) if renamed else '—'}。"
                f"旧值不删，见本行 git 历史与备份 {BAK.name}。")
    print(f"[曲目] {n0} → {len(out)} 首｜新增 {added}｜改名 {renamed}")
    old_note = str(ws.cell(r, note_c).value or "")
    if MARK in old_note:
        print("[=] 备注已含图源补录说明")
    else:
        print("[备注] 追加：", note_add[:160])
    if apply:
        ws.cell(r, song_c).value = "、".join(out)
        if MARK not in old_note:
            ws.cell(r, note_c).value = (old_note + "；" + note_add) if old_note else note_add
        if not BAK.exists():
            shutil.copy2(XLSX, BAK)
            print("[备份]", BAK)
        wb.save(XLSX)
        print("[写入]", XLSX)
    else:
        print("（试算，加 --apply 落地）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
