# -*- coding: utf-8 -*-
"""《沉响与长歌》行 204 二次补齐：新增《岩缝中开出的花》（10-09 图源点歌段第4首）。

背景：粉丝整理图（权威）显示 10-09 点歌段实唱 4 首，我们旧记录只有 3 首，缺
《岩峰中开出的花》（旧写法）。归属：B站全程 BV1Mypv6zECb p28「21.傲日其愣-《岩缝中开出的花》」
→ 傲日其愣独唱（已入 data/performance_attribution.json 的 exclude_for_subject）。

做法：在行 204「演唱曲目」列**追加**该曲（幂等；已存在则跳过），并在「备注」追加登记说明。
用法：python -X utf8 project_b\\fix_performance_attribution_row204_song.py [--apply]
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

XLSX = Path(r"E:\wx\index_records\王晰演出活动.xlsx")
BACKUP = XLSX.with_name("王晰演出活动.备份20261013b_岩缝补录.xlsx")
ROW = 204
NEW_SONG = "岩缝中开出的花"
MARK = "【补录 20261013】"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    seq_col, note_col, song_col = (hdr.index("序号") + 1, hdr.index("备注") + 1,
                                   hdr.index("演唱曲目") + 1)
    target = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, seq_col).value == ROW), None)
    if target is None:
        print("[X] 找不到序号 %d 的行" % ROW)
        return 1
    cur = str(ws.cell(target, song_col).value or "")
    if NEW_SONG in cur:
        print("[=] 演唱曲目已含《%s》，跳过" % NEW_SONG)
    else:
        newval = (cur.rstrip("、；;，, ") + "、" + NEW_SONG) if cur else NEW_SONG
        print("[曲目] %d 个分隔项 → %d 个" % (len(cur.split("、")), len(newval.split("、"))))
        if a.apply:
            ws.cell(target, song_col).value = newval
    note = str(ws.cell(target, note_col).value or "")
    if MARK in note:
        print("[=] 备注已含补录说明，跳过")
    else:
        add = (MARK + "粉丝整理图（权威）显示 10-09 点歌段**实唱 4 首**，旧记录缺《%s》；"
               "归属：B站全程 BV1Mypv6zECb p28「21.傲日其愣-《%s》」→ **傲日其愣独唱**，"
               "已计入同场「傲日独唱」剔除清单（data\\performance_attribution.json）。"
               "另：同场《南屏晚钟》《一生中最爱》系 10-10 点歌段，10-09 未唱，"
               "本行保留但标注（10-10 已由 --append-setlist 登记）。" % (NEW_SONG, NEW_SONG))
        print("[备注] 追加：", add[:220])
        if a.apply:
            ws.cell(target, note_col).value = (note + "；" + add) if note else add
    if a.apply:
        if not BACKUP.exists():
            shutil.copy2(XLSX, BACKUP)
            print("[备份]", BACKUP)
        wb.save(XLSX)
        print("[写入] xlsx")
    else:
        print("（试算，未写盘；加 --apply 落地）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
