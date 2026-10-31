"""phrase_bank 全量重建（v4）：mystic / interview / song / talk / lyrics 五类，均从源头重算"""
import glob, json, os, random, re
import openpyxl
from collections import Counter

ROOT = r"D:\wx409.github.io"
TALK_DIR = r"E:\wx\论文素材_王晰作传\现场talk"
SRC = r"E:\wx\骰子周边20260608\GEO\GEO"
BANK = os.path.join(ROOT, "data", "phrase_bank.json")
LYR = os.path.join(ROOT, "tavern", "lyrics_fragments.json")
DICE = os.path.join(ROOT, "data", "dice_sets.json")

items = []

# 1) 玄学（整句引语）
ws = openpyxl.load_workbook(os.path.join(SRC, r"vs\易经卦象解析表.xlsx"), read_only=True, data_only=True).active
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not r or not r[0]:
        continue
    q = r[3] if len(r) > 3 else None
    if not q:
        continue
    t = str(q).strip().strip("「」").strip()
    if 10 <= len(t) <= 80:
        items.append({"text": t, "type": "mystic", "source": "易经卦象解析表 第%d行｜%s（%s）" % (i, r[0], r[1] if len(r) > 1 else ""),
                      "curated": False, "rights": "娱乐向", "mood": "mystic", "verified": True})

# 2) 采访（标题 / 备注）
ws = openpyxl.load_workbook(os.path.join(SRC, "王晰采访记录汇编.xlsx"), read_only=True, data_only=True).active
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not r or not r[0]:
        continue
    name = r[1] if len(r) > 1 else None
    note = r[4] if len(r) > 4 else None
    if name:
        items.append({"text": str(name).strip(), "type": "interview", "source": "采访记录汇编 第%d行" % i,
                      "curated": False, "rights": "free", "mood": "title", "verified": True})
    if note and len(str(note).strip()) >= 10:
        items.append({"text": str(note).strip(), "type": "interview", "source": "采访记录汇编 第%d行" % i,
                      "curated": False, "rights": "free", "mood": "note", "verified": False})

# 3) 歌曲（歌名＋官方说明）
ws = openpyxl.load_workbook(os.path.join(SRC, "王晰歌曲信息汇总.xlsx"), read_only=True, data_only=True).active
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not r or len(r) < 7 or not r[2]:
        continue
    song, content = r[2], r[6]
    if content and str(content).strip():
        items.append({"text": "%s：%s" % (song, str(content).strip()), "type": "song",
                      "source": "王晰歌曲信息汇总 第%d行" % i, "curated": False,
                      "rights": "free", "mood": "song", "verified": True})

# 4) 歌词片段
ly = json.load(open(LYR, encoding="utf-8"))
for song, v in ly.items():
    if song == "_meta" or not isinstance(v, dict):
        continue
    for fr in v.get("fragments") or []:
        t = (fr.get("text") or "").strip()
        if t:
            items.append({"text": t, "type": "lyrics", "source": "歌词《%s》" % song, "curated": False,
                          "rights": "仅内部", "mood": "/".join(fr.get("moods") or []), "verified": True})

# 5) 现场 talk（ASR 短句，折叠重复字）
MARK = re.compile(r"(?:嗯+|然后|所以|但是|其实|我觉得|就是说|你知道吗|对吧|是吧|真的|谢谢|感谢|因为|就是|对啊|你看|我跟你说|那个时候)")
BAD_END = set("的但在我你他就也还 把被让给对和跟从到一是在那这说有着了过非常特别真")
EMO = ["谢谢", "感谢", "开心", "希望", "温暖", "幸福", "喜欢", "珍惜", "记得", "陪伴", "勇敢", "相信", "感动", "努力"]

def clean(t):
    t = re.sub(r"(.)\1{2,}", r"\1\1", t)
    t = re.sub(r"^这个|^那个", "", t)
    t = re.sub(r"(非常|特别|真的)$", "", t)
    return t.strip()

seen = set()
for f in sorted(glob.glob(os.path.join(TALK_DIR, "**", "*.txt"), recursive=True)):
    rel = os.path.relpath(f, TALK_DIR)
    city = rel.split(os.sep)[0] if os.sep in rel else ""
    for p in MARK.split(open(f, encoding="utf-8", errors="ignore").read()):
        t = clean(re.sub(r"[\s\dA-Za-z，。！？、；：\"'（）()…—\-]+", "", p or ""))
        if not (12 <= len(t) <= 28) or t[-1] in BAD_END:
            continue
        if not re.search(r"[我你咱]", t):
            continue
        if len(re.findall(r"[\u4e00-\u9fff]", t)) < len(t) * 0.95 or t in seen:
            continue
        seen.add(t)
        items.append({"text": t, "type": "talk", "source": "%s｜%s" % (city, os.path.basename(f)[:36]),
                      "curated": False, "rights": "free", "mood": "talk", "verified": False})

seen, uniq = set(), []
for it in items:
    if it["text"] in seen:
        continue
    seen.add(it["text"]); uniq.append(it)

stat = Counter(it["type"] for it in uniq)
json.dump({"generated": "2026-10-31", "count": len(uniq), "by_type": dict(stat),
           "note": "talk=现场转写短句（未校对 verified=false）；lyrics=歌词片段（仅内部）；song=歌名+官方说明；mystic=卦象整句；interview=标题(mood=title)/备注(mood=note)。",
           "items": uniq}, open(BANK, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("phrase_bank:", len(uniq), dict(stat))

random.seed(20261031)
def pick(types, n=6, maxlen=48, mood=None):
    pool = [x["text"] for x in uniq if x["type"] in types and 10 <= len(x["text"]) <= maxlen
            and (mood is None or x.get("mood") == mood)]
    random.shuffle(pool)
    return pool[:n]

talks = sorted([x["text"] for x in uniq if x["type"] == "talk"],
               key=lambda t: (sum(2 for k in EMO if k in t), 1 if 16 <= len(t) <= 26 else 0), reverse=True)

sets = {
    "emotion": {"name": "情绪骰", "disclaimer": "现场说话原句（语音转写、未校对，可能有断句误差）", "faces": talks[:6]},
    "mystic": {"name": "玄学骰", "disclaimer": "娱乐向，非占卜；引语出自站内整理的卦象解析表", "faces": pick(("mystic",), maxlen=80)},
    "lyrics": {"name": "歌词骰", "disclaimer": "短句引用，版权归原作者；不整首、不成集", "faces": pick(("lyrics",), maxlen=40)},
    "song": {"name": "歌名骰", "disclaimer": "歌名＋官方一句说明", "faces": pick(("song",), maxlen=60)},
    "interview": {"name": "采访骰", "disclaimer": "出自采访记录汇编；部分为标题或备注，非逐字引语", "faces": pick(("interview",), maxlen=60, mood="note")},
}
for v in sets.values():
    v["face_count"] = len(v["faces"])
json.dump({"generated": "2026-10-31", "sets": sets, "pending": []},
          open(DICE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for k, v in sets.items():
    print("==", v["name"], v["face_count"])
    for f in v["faces"]:
        print("   ", f)
